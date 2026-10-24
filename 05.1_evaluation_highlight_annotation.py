import os
import pandas as pd
import openpyxl
from openpyxl.styles import Font
from collections import defaultdict
import re


def compute_scores(TP, TN, FP, FN, decimals=None):
    """计算各项分数"""
    if decimals is None:
        decimals = KEEP_DECIMALS

    scores = {}
    scores['accuracy'] = round((TP + TN) / (TP + TN + FP + FN), decimals)
    scores['precision'] = round(TP / (TP + FP), decimals)
    scores['recall'] = round(TP / (TP + FN), decimals)
    p, r = scores['precision'], scores['recall']
    scores['f1'] = round(2 * TP / (2 * TP + FP + FN), decimals) if (p + r) > 0 else 0.0

    return scores


def count_excel_colors(file_path, exclude_cols=None, decimals=None):
    """统计Excel中的颜色数量"""
    if decimals is None:
        decimals = KEEP_DECIMALS

    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    # 颜色对应的类型
    colors = {
        '00B050': 'TP',  # 绿色
        'FFFF00': 'TN',  # 黄色
        '0070C0': 'FP',  # 蓝色
        '7030A0': 'FN'  # 紫色
    }

    # 读取第一行作为列名
    first_row = 1
    col_names = {}

    for col_idx, cell in enumerate(ws[first_row], 1):
        letter = cell.column_letter
        col_names[letter] = str(cell.value) if cell.value else f"列{letter}"

    # 设置要跳过的列
    if exclude_cols is None:
        exclude_cols = [list(col_names.values())[0]]

    # 找到要跳过的列字母
    skip_letters = []
    for col_name in exclude_cols:
        for letter, name in col_names.items():
            if name == col_name:
                skip_letters.append(letter)
                break

    # 按列统计颜色
    col_counts = defaultdict(lambda: {'TP': 0, 'TN': 0, 'FP': 0, 'FN': 0})
    all_counts = {'TP': 0, 'TN': 0, 'FP': 0, 'FN': 0}

    # 从第二行开始统计
    for row in ws.iter_rows(min_row=2, max_col=ws.max_column, max_row=ws.max_row):
        for cell in row:
            letter = cell.column_letter

            if letter in skip_letters:
                continue

            if cell.fill and cell.fill.start_color:
                rgb = str(cell.fill.start_color.rgb or '').upper()

                if rgb and len(rgb) >= 8:
                    hex_color = rgb[2:]

                    if len(hex_color) > 6:
                        hex_color = hex_color[:6]

                    if hex_color in colors:
                        color_type = colors[hex_color]

                        col_name = col_names[letter]
                        col_counts[col_name][color_type] += 1
                        all_counts[color_type] += 1

    # 计算各列分数
    result = {}
    for col_name, counts in col_counts.items():
        result[col_name] = {
            'counts': counts.copy(),
            'scores': compute_scores(**counts, decimals=decimals)
        }

    # 计算总计
    if any(all_counts.values()):
        result['总计'] = {
            'counts': all_counts.copy(),
            'scores': compute_scores(**all_counts, decimals=decimals)
        }

    wb.close()
    return result


def get_model_name(filename):
    """从文件名提取模型名"""
    pattern = r"md_(.+?)_color\.xlsx"
    match = re.search(pattern, filename, re.IGNORECASE)
    if match:
        return match.group(1)

    pattern2 = r"(.+?)_color\.xlsx"
    match = re.search(pattern2, filename, re.IGNORECASE)
    if match:
        return match.group(1)

    return os.path.splitext(filename)[0]


def process_folder(folder_path, exclude_cols=None, decimals=None):
    """处理整个文件夹"""
    if decimals is None:
        decimals = KEEP_DECIMALS

    # 找到所有Excel文件
    excel_files = []
    for f in os.listdir(folder_path):
        if f.lower().endswith('.xlsx'):
            excel_files.append(f)

    if not excel_files:
        print(f"文件夹 {folder_path} 中没有找到符合条件的Excel文件")
        return {}

    print(f"找到 {len(excel_files)} 个文件:")
    for f in excel_files:
        model_name = get_model_name(f)
        print(f"  - {f} -> {model_name}")

    # 逐个处理
    all_data = {}
    for f in excel_files:
        file_path = os.path.join(folder_path, f)
        model_name = get_model_name(f)
        try:
            wb = openpyxl.load_workbook(file_path)
            ws = wb.active
            first_col = str(ws[1][0].value) if ws[1][0].value else "列A"
            wb.close()

            # 设置跳过哪些列
            if exclude_cols is None:
                skip_cols = [first_col]
            else:
                skip_cols = exclude_cols

            # 统计颜色
            data = count_excel_colors(
                file_path,
                exclude_cols=skip_cols,
                decimals=decimals
            )

            all_data[model_name] = data

        except Exception as err:
            print(f"  出错: {err}")

    return all_data


def make_data_tables(all_data, decimals=None):
    """制作数据表，行=列名，列=模型"""
    if decimals is None:
        decimals = KEEP_DECIMALS

    # 收集所有列名（跳过'总计'）
    all_cols = set()
    for model_data in all_data.values():
        for col_name in model_data.keys():
            if col_name != '总计':
                all_cols.add(col_name)

    all_cols = sorted(list(all_cols))
    print(all_cols)

    # 收集所有模型名
    all_models = list(all_data.keys())

    # 四个表的字典
    tables = {}

    for table_name in ['F1', 'Precision', 'Recall', 'Accuracy']:
        # 新建空表
        df = pd.DataFrame(index=all_cols, columns=all_models)

        # 填数据
        for col_name in all_cols:
            for model_name in all_models:
                if col_name in all_data[model_name]:
                    scores = all_data[model_name][col_name]['scores']
                    if table_name == 'F1':
                        val = scores['f1']
                    elif table_name == 'Precision':
                        val = scores['precision']
                    elif table_name == 'Recall':
                        val = scores['recall']
                    elif table_name == 'Accuracy':
                        val = scores['accuracy']
                    df.loc[col_name, model_name] = round(val, decimals)
                else:
                    df.loc[col_name, model_name] = None

        # 计算每个模型Average
        # 使用总TP/TN/FP/FN计算Average（micro average）
        averages = {}

        for model_name in all_models:

            total_counts = all_data[model_name]['总计']['counts']

            total_scores = compute_scores(
                TP=total_counts['TP'],
                TN=total_counts['TN'],
                FP=total_counts['FP'],
                FN=total_counts['FN'],
                decimals=decimals
            )

            if table_name == 'F1':
                avg = total_scores['f1']

            elif table_name == 'Precision':
                avg = total_scores['precision']

            elif table_name == 'Recall':
                avg = total_scores['recall']

            elif table_name == 'Accuracy':
                avg = total_scores['accuracy']

            averages[model_name] = avg

        # 添加Average行
        df.loc['Average'] = pd.Series(averages)

        tables[table_name] = df

    return tables


def write_to_excel(tables, save_path, decimals=None):
    """写入Excel文件"""
    if decimals is None:
        decimals = KEEP_DECIMALS

    # 设置数字格式
    if decimals == 0:
        num_fmt = '0'
    elif decimals == 1:
        num_fmt = '0.0'
    elif decimals == 2:
        num_fmt = '0.00'
    elif decimals == 3:
        num_fmt = '0.000'
    elif decimals == 4:
        num_fmt = '0.0000'
    else:
        num_fmt = f'0.{"0" * decimals}'

    # 创建汇总表
    all_df = create_all_summary_table(tables, decimals)

    # 写入Excel
    with pd.ExcelWriter(save_path, engine='openpyxl') as writer:
        # 先写入'all'表
        all_df.to_excel(writer, index=False, sheet_name='all')

        # 获取all表的工作表
        ws_all = writer.sheets['all']

        # 设置数字格式
        for r in range(2, len(all_df) + 2):
            for c in range(2, len(all_df.columns) + 1):
                cell = ws_all.cell(row=r, column=c)

                if isinstance(cell.value, (int, float)):

                    # F1 / Precision / Recall / Accuracy
                    if c <= 5:
                        cell.number_format = num_fmt

                    # TP / TN / FP / FN
                    else:
                        cell.number_format = '0'

        # 设置列宽
        for col in ws_all.columns:
            max_len = 0
            letter = col[0].column_letter
            for cell in col:
                try:
                    if cell.value and len(str(cell.value)) > max_len:
                        max_len = len(str(cell.value))
                except:
                    pass
            width = min(max_len + 2, 30)
            ws_all.column_dimensions[letter].width = width

        # 写入其他表
        for sheet_name, df in tables.items():
            # 设置行标题
            df.index.name = "列名"
            df.to_excel(writer, index=True, sheet_name=sheet_name)

            # 获取工作表
            ws = writer.sheets[sheet_name]

            # 设置数字格式
            for r in range(2, len(df) + 2):
                for c in range(2, len(df.columns) + 2):
                    cell = ws.cell(row=r, column=c)
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = num_fmt

            # Average行加粗
            last_row = len(df)
            for c in range(1, len(df.columns) + 2):
                cell = ws.cell(row=last_row + 1, column=c)
                cell.font = Font(bold=True)

            # 调整列宽
            for col in ws.columns:
                max_len = 0
                letter = col[0].column_letter
                for cell in col:
                    try:
                        if cell.value and len(str(cell.value)) > max_len:
                            max_len = len(str(cell.value))
                    except:
                        pass
                width = min(max_len + 2, 30)
                ws.column_dimensions[letter].width = width

    return tables, all_df


def create_all_summary_table(tables, decimals=None):
    """创建汇总所有模型平均值的'all'表"""
    if decimals is None:
        decimals = KEEP_DECIMALS

    # 获取所有模型名
    all_models = list(tables['F1'].columns)

    # 创建汇总DataFrame
    summary_data = []

    for model_name in all_models:
        counts = all_data[model_name]['总计']['counts']
        tp = counts['TP']
        tn = counts['TN']
        fp = counts['FP']
        fn = counts['FN']

        # 从各个表中获取对应模型的Average值
        f1_avg = tables['F1'].loc['Average', model_name]
        precision_avg = tables['Precision'].loc['Average', model_name]
        recall_avg = tables['Recall'].loc['Average', model_name]
        accuracy_avg = tables['Accuracy'].loc['Average', model_name]

        summary_data.append({
            'model': model_name,
            'Precision': precision_avg,
            'Recall': recall_avg,
            'F1-score': f1_avg,
            'Accuracy': accuracy_avg,
            'TP': tp,
            'TN': tn,
            'FP': fp,
            'FN': fn,
        })

    # 创建DataFrame
    summary_df = pd.DataFrame(summary_data)

    return summary_df


if __name__ == "__main__":
    if __name__ == "__main__":
        a = 'raw-50-1'
        input_folder = f"E:\\dm\\model-tag-{a}\\model-color"
        output_file = f'E:\\dm\\model-tag-{a}\\model-score.xlsx'
        KEEP_DECIMALS = 2  # 这里设置小数位数

        # 处理文件
        print(f"处理文件夹: {input_folder}")
        all_data = process_folder(
            folder_path=input_folder,
            exclude_cols=None,
            decimals=KEEP_DECIMALS
        )

        tables = make_data_tables(all_data, decimals=KEEP_DECIMALS)
        print(f"\n保存到: {output_file}")
        tables, all_df = write_to_excel(tables, output_file, decimals=KEEP_DECIMALS)  # 接收返回值
        os.startfile(output_file)